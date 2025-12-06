import re
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


class ElementBuilder:
    """Класс для построения элементов документа по ГОСТ"""

    def __init__(self, counters_manager):
        self.counters_manager = counters_manager

    def add_normal_text_gost(self, source_paragraph, target_doc):
        """Добавление обычного текста по ГОСТ"""
        if not source_paragraph.text.strip():
            return None

        paragraph = target_doc.add_paragraph()

        for run in source_paragraph.runs:
            new_run = paragraph.add_run(run.text)
            new_run.font.name = 'Times New Roman'
            new_run.font.size = Pt(14)
            new_run.font.color.rgb = RGBColor(0, 0, 0)
            new_run.bold = run.bold
            new_run.italic = run.italic
            new_run.underline = run.underline

        pf = paragraph.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.line_spacing = 1.5
        pf.first_line_indent = Inches(0.49)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.left_indent = Inches(0)
        pf.right_indent = Inches(0)

        return paragraph

    def add_image_placeholder_gost(self, doc, caption="Иллюстрация"):
        """Добавление изображения с нумерацией в пределах раздела"""
        figure_number = self.counters_manager.get_next_number('figure')

        placeholder = doc.add_paragraph()
        placeholder.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = placeholder.add_run("(место для изображения)")
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14)
        run.font.italic = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        placeholder.paragraph_format.space_after = Pt(6)

        caption_para = doc.add_paragraph()
        caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_run = caption_para.add_run(f"Рисунок {figure_number} — {caption}")
        caption_run.font.name = 'Times New Roman'
        caption_run.font.size = Pt(12)
        caption_run.font.bold = True
        caption_run.font.color.rgb = RGBColor(0, 0, 0)

        caption_para.paragraph_format.space_before = Pt(6)
        caption_para.paragraph_format.space_after = Pt(6)
        caption_para.paragraph_format.line_spacing = 1.0

        return figure_number

    def add_table_with_caption_gost(self, source_table, target_doc, caption="Название таблицы"):
        """Добавление таблицы с автоподбором и нумерацией в пределах раздела"""
        table_number = self.counters_manager.get_next_number('table')

        caption_para = target_doc.add_paragraph()
        caption_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        caption_run = caption_para.add_run(f"Таблица {table_number} — {caption}")
        caption_run.font.name = 'Times New Roman'
        caption_run.font.size = Pt(12)
        caption_run.font.italic = True
        caption_run.font.color.rgb = RGBColor(0, 0, 0)

        caption_para.paragraph_format.space_before = Pt(6)
        caption_para.paragraph_format.space_after = Pt(0)
        caption_para.paragraph_format.line_spacing = 1.0

        rows = len(source_table.rows)
        cols = len(source_table.columns)
        if rows > 0 and cols > 0:
            new_table = target_doc.add_table(rows=rows, cols=cols)
            new_table.style = 'Table Grid'
            new_table.autofit = True

            for row_idx, row in enumerate(source_table.rows):
                for col_idx, cell in enumerate(row.cells):
                    new_cell = new_table.cell(row_idx, col_idx)
                    new_cell.text = cell.text

                    for p in new_cell.paragraphs:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        for run in p.runs:
                            run.font.name = 'Times New Roman'
                            run.font.size = Pt(12)
                            run.font.color.rgb = RGBColor(0, 0, 0)

            for cell in new_table.rows[-1].cells:
                for p in cell.paragraphs:
                    if p.text.strip():
                        p.paragraph_format.space_after = Pt(6)

        return table_number

    def add_listing_caption_gost(self, target_doc, listing_number, caption_text=""):
        """Добавление подписи к листингу"""
        if caption_text and caption_text.strip():
            clean_caption = re.sub(r'^листинг\s*\d+(\.\d+)*\s*[—-]\s*', '', caption_text, flags=re.IGNORECASE)
            clean_caption = re.sub(r'^listing\s*\d+(\.\d+)*\s*[—-]\s*', '', clean_caption, flags=re.IGNORECASE)
            full_caption = f"Листинг {listing_number} — {clean_caption.strip()}"
        else:
            full_caption = f"Листинг {listing_number}"

        caption_para = target_doc.add_paragraph()

        caption_run = caption_para.add_run(full_caption)
        caption_run.font.name = 'Times New Roman'
        caption_run.font.size = Pt(12)
        caption_run.font.italic = True
        caption_run.font.color.rgb = RGBColor(0, 0, 0)

        pf = caption_para.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.left_indent = Inches(0)
        pf.right_indent = Inches(0)
        pf.first_line_indent = Inches(0)
        pf.space_before = Pt(6)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0

        return caption_para

    def add_listing_content_with_frame_gost(self, source_paragraph, target_doc):
        """Добавление листинга в рамке по ГОСТ"""
        # Создаем таблицу с одной ячейкой для рамки
        table = target_doc.add_table(rows=1, cols=1)
        table.autofit = False
        table.style = 'Table Grid'

        cell = table.cell(0, 0)
        cell.paragraphs[0].clear()

        listing_para = cell.add_paragraph()

        for run in source_paragraph.runs:
            new_run = listing_para.add_run(run.text)
            new_run.font.name = 'Courier New'
            new_run.font.size = Pt(10)
            new_run.font.color.rgb = RGBColor(0, 0, 0)
            new_run.bold = False
            new_run.italic = False
            new_run.underline = False

        pf = listing_para.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.left_indent = Inches(0)
        pf.right_indent = Inches(0)
        pf.first_line_indent = Inches(0)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0

        # Добавляем отступ после листинга
        target_doc.add_paragraph().paragraph_format.space_before = Pt(6)

        return listing_para

    def add_list_item_gost(self, source_paragraph, target_doc, list_hierarchy=0, list_type='bulleted', is_last=False,
                           group_id=None, num_id=None):
        """Добавление элемента списка"""
        text = source_paragraph.text.strip()
        clean_text = text

        # ПРОСТАЯ ЛОГИКА: если нумерованный - удаляем номер, если маркированный - удаляем маркер
        if list_type == 'numbered':
            # Удаляем номер в начале (например, "1. ", "10. ")
            match = re.match(r'^(\d+(?:\.\d+)*)\.\s*(.*)', text)
            if match:
                clean_text = match.group(2)
            else:
                clean_text = re.sub(r'^\d+(?:\.\d+)*\.\s*', '', text)
        elif list_type == 'bulleted':
            # Удаляем маркер в начале
            clean_text = re.sub(r'^[•\-—–]\s*', '', text)

        # Форматирование по ГОСТу
        if clean_text:
            if list_type == 'numbered':
                # Для нумерованных списков: первая буква заглавная, заканчивается точкой
                if clean_text and clean_text[0].isalpha():
                    clean_text = clean_text[0].upper() + clean_text[1:]

                if not clean_text.endswith(('.', '!', '?')):
                    clean_text = clean_text + '.'
            else:
                # Для маркированных списков: первая буква строчная
                if clean_text and clean_text[0].isalpha():
                    clean_text = clean_text[0].lower() + clean_text[1:]

                # Маркированные списки заканчиваются точкой или точкой с запятой
                if not clean_text.endswith(('.', '!', '?', ';')):
                    if is_last:
                        clean_text = clean_text + '.'
                    else:
                        clean_text = clean_text + ';'

        # Создаем параграф
        if list_type == 'numbered':
            paragraph = target_doc.add_paragraph()

            if num_id is None:
                pPr = paragraph._element.get_or_add_pPr()
                numPr = pPr.get_or_add_numPr()

                ilvl = OxmlElement('w:ilvl')
                ilvl.set(qn('w:val'), str(list_hierarchy))
                numPr.append(ilvl)

                numId = OxmlElement('w:numId')
                numId.set(qn('w:val'), str(1000 + (group_id if group_id is not None else 0)))
                numPr.append(numId)
            else:
                pPr = paragraph._element.get_or_add_pPr()
                numPr = pPr.get_or_add_numPr()

                ilvl = OxmlElement('w:ilvl')
                ilvl.set(qn('w:val'), str(list_hierarchy))
                numPr.append(ilvl)

                numId = OxmlElement('w:numId')
                numId.set(qn('w:val'), str(num_id))
                numPr.append(numId)
        else:
            paragraph = target_doc.add_paragraph(style='List Bullet')
            if list_hierarchy > 0:
                pPr = paragraph._element.get_or_add_pPr()
                numPr = pPr.get_or_add_numPr()
                ilvl = OxmlElement('w:ilvl')
                ilvl.set(qn('w:val'), str(list_hierarchy))
                numPr.append(ilvl)

                numId = OxmlElement('w:numId')
                numId.set(qn('w:val'), '2')
                numPr.append(numId)

        # Добавляем текст
        run = paragraph.add_run(clean_text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0, 0, 0)

        # Настраиваем форматирование параграфа
        pf = paragraph.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.line_spacing = 1.5

        marker_indent = 0.49 + (0.49 * list_hierarchy)
        text_indent = 0.89 + (0.49 * list_hierarchy)

        pf.left_indent = Inches(text_indent)
        pf.first_line_indent = Inches(-(text_indent - marker_indent))

        pf.space_before = Pt(0)
        pf.space_after = Pt(0)

        return paragraph