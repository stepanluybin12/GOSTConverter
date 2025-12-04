from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION


class StyleManager:
    """Класс для управления стилями ГОСТ"""

    @staticmethod
    def create_styles(doc):
        """Создание стилей ГОСТ"""
        try:
            styles_to_create = [
                ('Heading 1 GOST', WD_STYLE_TYPE.PARAGRAPH, doc.styles['Heading 1'], {
                    'name': 'Times New Roman',
                    'size': Pt(18),
                    'bold': True,
                    'color': RGBColor(0, 0, 0),
                    'alignment': WD_ALIGN_PARAGRAPH.JUSTIFY,
                    'space_before': Pt(0),
                    'space_after': Pt(12),
                    'line_spacing': 1.5,
                    'first_line_indent': Inches(0.5)
                }),
                ('Heading 2 GOST', WD_STYLE_TYPE.PARAGRAPH, doc.styles['Heading 2'], {
                    'name': 'Times New Roman',
                    'size': Pt(16),
                    'bold': True,
                    'color': RGBColor(0, 0, 0),
                    'alignment': WD_ALIGN_PARAGRAPH.JUSTIFY,
                    'space_before': Pt(24),
                    'space_after': Pt(12),
                    'line_spacing': 1.5,
                    'first_line_indent': Inches(0.5),
                }),
                ('Heading 3 GOST', WD_STYLE_TYPE.PARAGRAPH, doc.styles['Heading 3'], {
                    'name': 'Times New Roman',
                    'size': Pt(14),
                    'bold': True,
                    'color': RGBColor(0, 0, 0),
                    'alignment': WD_ALIGN_PARAGRAPH.JUSTIFY,
                    'space_before': Pt(24),
                    'space_after': Pt(12),
                    'line_spacing': 1.5,
                    'first_line_indent': Inches(0.5),
                }),
            ]

            for style_name, style_type, base_style, params in styles_to_create:
                if style_name not in [s.name for s in doc.styles]:
                    style = doc.styles.add_style(style_name, style_type)
                    style.base_style = base_style
                    font = style.font
                    font.name = params['name']
                    font.size = params['size']
                    font.bold = params['bold']
                    font.color.rgb = RGBColor(0, 0, 0)
                    pf = style.paragraph_format
                    pf.alignment = params['alignment']
                    pf.space_before = params['space_before']
                    pf.space_after = params['space_after']
                    pf.line_spacing = params['line_spacing']
                    if 'first_line_indent' in params:
                        pf.first_line_indent = params['first_line_indent']

        except Exception as e:
            print(f"Не удалось создать стили: {e}")

    @staticmethod
    def create_empty_section_without_footer(doc):
        """Создает новый раздел без нижнего колонтитула"""
        # Добавляем новый раздел
        doc.add_section(WD_SECTION.NEW_PAGE)
        new_section = doc.sections[-1]

        # Настраиваем поля
        new_section.top_margin = Inches(0.79)
        new_section.bottom_margin = Inches(0.79)
        new_section.left_margin = Inches(1.18)
        new_section.right_margin = Inches(0.39)

        # Очищаем нижний колонтитул
        StyleManager.clear_footer(new_section)

        return new_section

    @staticmethod
    def clear_footer(section):
        """Очищает нижний колонтитул раздела"""
        footer = section.footer

        # Удаляем все параграфы из нижнего колонтитула
        for paragraph in footer.paragraphs:
            p = paragraph._element
            p.getparent().remove(p)

    @staticmethod
    def create_page_number_footer(section):
        """Создание нижнего колонтитула с нумерацией страниц"""
        footer = section.footer

        # Очищаем существующий нижний колонтитул
        StyleManager.clear_footer(section)

        paragraph = footer.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run()

        # Добавляем поле PAGE
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')

        instrText = OxmlElement('w:instrText')
        instrText.text = "PAGE"
        instrText.set(qn('xml:space'), 'preserve')

        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'end')

        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)

        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)

        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)

        # Устанавливаем расстояние от нижнего колонтитула
        section.footer_distance = Inches(0.3)

    @staticmethod
    def set_section_starting_page_number(section, start_number):
        """Устанавливает начальный номер страницы для раздела"""
        sectPr = section._sectPr

        # Находим или создаем элемент pgNumType
        pgNumType = sectPr.find(qn('w:pgNumType'))
        if pgNumType is None:
            pgNumType = OxmlElement('w:pgNumType')
            sectPr.append(pgNumType)

        # Устанавливаем начальный номер страницы
        pgNumType.set(qn('w:start'), str(start_number))